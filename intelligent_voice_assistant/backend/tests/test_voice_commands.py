from types import SimpleNamespace

import openpyxl

from backend.api import routes
from backend.services import excel_service
from backend.services.parser_service import parse_voice_command


def test_voice_command_accepts_short_roll_number_for_prefixed_excel_id():
    assert excel_service._roll_matches("23BSIT-01", 1)
    assert excel_service._roll_matches("24BSIT-02", "2")
    assert not excel_service._roll_matches("23BSIT-11", 1)


def test_voice_parser_accepts_spoken_short_roll_and_chained_prefixed_rolls():
    result = parse_voice_command("roll no one marks two")
    assert result["roll_no"] == 1
    assert result["marks"] == 2

    result = parse_voice_command(
        "roll no 23BSIT-01 marks 2 and roll no 24BSIT-02 marks 4"
    )
    assert result["multi"] is True
    assert [item["roll_no"] for item in result["commands"]] == [
        "23BSIT-01",
        "24BSIT-02",
    ]


def test_voice_command_updates_multiple_columns_without_locking(monkeypatch):
    calls = []

    def fake_update_marks(roll_no, marks):
        calls.append((routes.get_current_column(), roll_no, marks))
        return True, "Marks Updated Successfully."

    monkeypatch.setattr(routes, "update_marks", fake_update_marks)
    monkeypatch.setattr(routes, "get_current_column", lambda: None)
    monkeypatch.setattr(routes, "set_current_column", lambda column: calls.append(("selected", column, None)))

    response = routes.voice_command(
        SimpleNamespace(text="roll no 1 assignment marks 20 test marks 30 midterm marks 40")
    )

    assert response["updated_count"] == 3
    assert [(roll, marks) for kind, roll, marks in calls if kind != "selected"] == [
        (1, 20),
        (1, 30),
        (1, 40),
    ]


def test_voice_column_only_command_does_not_lock_column(monkeypatch):
    monkeypatch.setattr(routes, "get_current_column", lambda: "assignment")
    response = routes.voice_command(SimpleNamespace(text="select final"))

    assert response["voice_column_ignored"] is True
    assert "roll number" in response["message"]


def test_voice_command_accepts_presentation_before_roll(monkeypatch):
    calls = []

    monkeypatch.setattr(routes, "get_current_column", lambda: None)
    monkeypatch.setattr(routes, "set_current_column", lambda column: None)
    monkeypatch.setattr(
        routes,
        "update_marks",
        lambda roll_no, marks: calls.append((roll_no, marks)) or (True, "Marks Updated Successfully."),
    )

    response = routes.voice_command(
        SimpleNamespace(text="presentation roll no 1 marks 4")
    )

    assert response["updated"] is True
    assert response["results"][0]["column"] == "presentation"
    assert calls == [(1, 4)]


def test_final_and_final_term_are_same_assessment():
    final = parse_voice_command("final roll no 1 marks 35")
    final_term = parse_voice_command("final term roll no 1 marks 35")

    assert final["column"] == "final"
    assert final_term["column"] == "final"


def test_common_presentation_speech_misspelling_is_normalized():
    result = parse_voice_command("presenation roll no 1 marks 4")
    assert result["column"] == "presentation"
    assert result["marks"] == 4


def test_voice_command_updates_multiple_roll_entries_in_one_sentence(monkeypatch):
    calls = []

    def fake_update_marks(roll_no, marks):
        calls.append((roll_no, marks))
        return True, f"Marks Updated Successfully for roll {roll_no}."

    monkeypatch.setattr(routes, "update_marks", fake_update_marks)
    monkeypatch.setattr(routes, "get_current_column", lambda: "assignment")

    payload = SimpleNamespace(text="roll no 1 marks 80 and roll no 2 marks 90")
    response = routes.voice_command(payload)

    assert response["multi"] is True
    assert response["updated_count"] == 2
    assert calls == [(1, 80), (2, 90)]
    assert response["message"] == "Updated 2 marks entries."


def test_voice_command_splits_back_to_back_roll_entries(monkeypatch):
    calls = []

    def fake_update_marks(roll_no, marks):
        calls.append((roll_no, marks))
        return True, f"Marks Updated Successfully for roll {roll_no}."

    monkeypatch.setattr(routes, "update_marks", fake_update_marks)
    monkeypatch.setattr(routes, "get_current_column", lambda: "assignment")

    payload = SimpleNamespace(text="roll no 1 marks 2 roll no 2 marks 3")
    response = routes.voice_command(payload)

    assert response["multi"] is True
    assert response["updated_count"] == 2
    assert calls == [(1, 2), (2, 3)]
    assert response["message"] == "Updated 2 marks entries."


def test_voice_command_updates_row_wise_marks_for_one_roll(monkeypatch):
    calls = []

    def fake_update_marks(roll_no, marks):
        calls.append((routes.get_current_column(), roll_no, marks))
        return True, "Marks Updated Successfully."

    monkeypatch.setattr(routes, "update_marks", fake_update_marks)
    monkeypatch.setattr(routes, "get_current_column", lambda: "final")
    monkeypatch.setattr(routes, "set_current_column", lambda column: setattr(routes, "_test_column", column))
    routes._test_column = "final"
    monkeypatch.setattr(routes, "get_current_column", lambda: routes._test_column)

    response = routes.voice_command(
        SimpleNamespace(text="roll number 1 assignment 8 midterm 25 final 35")
    )

    assert response["updated_count"] == 3
    assert calls == [
        ("assignment", 1, 8),
        ("midterm", 1, 25),
        ("final", 1, 35),
    ]


def test_voice_command_understands_spoken_number_words(monkeypatch):
    calls = []

    def fake_update_marks(roll_no, marks):
        calls.append((roll_no, marks))
        return True, f"Marks Updated Successfully for roll {roll_no}."

    monkeypatch.setattr(routes, "update_marks", fake_update_marks)
    monkeypatch.setattr(routes, "get_current_column", lambda: "assignment")

    payload = SimpleNamespace(text="roll number 1 marks eighty five")
    response = routes.voice_command(payload)

    assert response["updated"] is True
    assert calls == [(1, 85)]
    assert response["message"] == "Marks Updated Successfully."


def test_set_column_route_persists_selected_column(monkeypatch):
    captured = {}

    def fake_set_current_column(column_name):
        captured["column"] = column_name

    monkeypatch.setattr(routes, "set_current_column", fake_set_current_column)

    response = routes.set_column(SimpleNamespace(column="Assignment"))

    assert response["selected_column"] == "assignment"
    assert captured["column"] == "assignment"


def test_voice_command_accepts_clear_natural_phrases(monkeypatch):
    calls = []

    def fake_update_marks(roll_no, marks):
        calls.append((roll_no, marks))
        return True, f"Marks Updated Successfully for roll {roll_no}."

    monkeypatch.setattr(routes, "update_marks", fake_update_marks)
    monkeypatch.setattr(routes, "get_current_column", lambda: "assignment")

    response = routes.voice_command(SimpleNamespace(text="assignment me roll number 1 marks 85"))
    assert response["selected_column"] == "assignment"
    assert response["multi"] is True
    assert calls == [(1, 85)]

    response = routes.voice_command(SimpleNamespace(text="marks 85 roll number 1"))
    assert response["selected_column"] == "assignment"
    assert response["updated"] is True
    assert response["results"][0]["roll_no"] == 1
    assert response["results"][0]["marks"] == 85


def test_voice_command_shows_range_error_for_invalid_marks(monkeypatch):
    monkeypatch.setattr(routes, "get_current_column", lambda: "midterm")
    monkeypatch.setattr(
        routes,
        "update_marks",
        lambda roll_no, marks: (False, "Midterm marks must be between 1 and 30."),
    )

    response = routes.voice_command(
        SimpleNamespace(text="midterm roll number 1 marks 35")
    )

    assert response["updated"] is False
    assert response["message"] == "Midterm marks must be between 1 and 30."


def test_update_marks_rejects_invalid_test_marks(tmp_path, monkeypatch):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Roll_No", "Test"])
    workbook.save(tmp_path / "test.xlsx")

    monkeypatch.setattr(excel_service, "EXCEL_FILE", str(tmp_path / "test.xlsx"))
    monkeypatch.setattr(excel_service, "get_current_column", lambda: "test")

    success, message = excel_service.update_marks(1, 11)

    assert success is False
    assert message == "Test marks must be between 1 and 10."


def test_update_marks_creates_row_when_roll_does_not_exist(tmp_path, monkeypatch):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "Roll_No"
    sheet["B1"] = "Assignment"
    workbook.save(tmp_path / "test.xlsx")

    monkeypatch.setattr(excel_service, "EXCEL_FILE", str(tmp_path / "test.xlsx"))
    monkeypatch.setattr(excel_service, "get_current_column", lambda: "assignment")

    success, message = excel_service.update_marks(90, 8)

    assert success is True
    assert "New row created" in message

    updated = openpyxl.load_workbook(tmp_path / "test.xlsx", data_only=True)
    row_values = list(updated.active.iter_rows(min_row=1, max_row=2, values_only=True))
    assert row_values[1][0] == 90
    assert row_values[1][1] == 8


def test_update_marks_creates_missing_presentation_column(tmp_path, monkeypatch):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Roll_No", "Test"])
    sheet.append(["23BSIT-01", 0])
    workbook.save(tmp_path / "test.xlsx")

    monkeypatch.setattr(excel_service, "EXCEL_FILE", str(tmp_path / "test.xlsx"))
    monkeypatch.setattr(excel_service, "get_current_column", lambda: "presentation")

    success, message = excel_service.update_marks(1, 4)

    assert success is True
    assert "updated" in message.lower()
    updated = openpyxl.load_workbook(tmp_path / "test.xlsx", data_only=True)
    assert [cell.value for cell in updated.active[1]] == ["Roll_No", "Test", "Presentation"]
    assert updated.active.cell(row=2, column=3).value == 4


def test_update_marks_rejects_marks_outside_selected_column_range(tmp_path, monkeypatch):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["Roll_No", "Assignment", "Midterm", "Final"])
    workbook.save(tmp_path / "test.xlsx")

    monkeypatch.setattr(excel_service, "EXCEL_FILE", str(tmp_path / "test.xlsx"))

    for column, marks, expected_range in [
        ("assignment", 11, "1 and 10"),
        ("midterm", 31, "1 and 30"),
        ("final", 51, "1 and 50"),
    ]:
        monkeypatch.setattr(excel_service, "get_current_column", lambda column=column: column)
        success, message = excel_service.update_marks(1, marks)

        assert success is False
        assert expected_range in message
